"""Build the fixed 10-slide Kravii presentation deck with native shapes only."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("kravii-deck.pptx")
W, H = Inches(13.333), Inches(7.5)

# Design plan: a white, technical briefing canvas. Navy establishes trust; a single
# electric teal accent denotes guarded actions and on-chain enforcement. The signature
# is the recurring teal "contract boundary" line used in architecture and flows.
NAVY = RGBColor(12, 30, 53)
TEAL = RGBColor(0, 143, 151)
INK = RGBColor(20, 31, 43)
MUTED = RGBColor(83, 98, 112)
LINE = RGBColor(207, 216, 222)
PALE = RGBColor(240, 244, 246)
GREY = RGBColor(210, 215, 219)
WHITE = RGBColor(255, 255, 255)


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def text(slide, value, x, y, w, h, size=20, color=INK, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear(); frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Inches(0)
    frame.vertical_anchor = valign
    para = frame.paragraphs[0]
    para.alignment = align
    para.text = value
    para.font.name = font; para.font.size = Pt(size); para.font.bold = bold; para.font.color.rgb = color
    para.space_after = Pt(0)
    return shape


def label(slide, value, x, y, w, h, size=12, color=TEAL, align=PP_ALIGN.LEFT):
    return text(slide, value.upper(), x, y, w, h, size, color, True, "Aptos Display", align)


def title(slide, value, kicker=None, size=36):
    if kicker:
        label(slide, kicker, 0.67, 0.36, 9.8, 0.25)
    text(slide, value, 0.67, 0.66, 12.0, 1.02, size, NAVY, True, "Aptos Display")
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.67), Inches(1.72), Inches(1.0), Inches(0.055))
    line.fill.solid(); line.fill.fore_color.rgb = TEAL; line.line.fill.background()


def footer(slide, number):
    text(slide, "KRAVII  /  FINTECHATHON 2026", 0.67, 7.12, 4.0, 0.16, 8, MUTED, True)
    text(slide, f"{number:02d}", 12.18, 7.08, 0.48, 0.2, 9, MUTED, True, align=PP_ALIGN.RIGHT)


def note(slide, value):
    frame = slide.notes_slide.notes_text_frame
    frame.text = value


def connector(slide, x1, y1, x2, y2, color=TEAL, arrow=True):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color; conn.line.width = Pt(1.5)
    if arrow:
        conn.line.end_arrowhead = True
    return conn


def card(slide, heading, body, x, y, w, h, accent=False):
    s = box(slide, x, y, w, h, WHITE, TEAL if accent else LINE, True)
    if accent:
        s.line.width = Pt(2)
    text(slide, heading, x + .18, y + .17, w - .36, .28, 17, NAVY, True, "Aptos Display")
    text(slide, body, x + .18, y + .56, w - .36, h - .68, 14, MUTED)
    return s


NOTES = [
"""Hello, we are team Kravii, formerly called Agentic Wallet. We built a wallet that an AI agent can run, and that still cannot lose your money, even if the agent is fooled. This is our entry for Topic F, Agent Wallet. In the next ten minutes: the problem, how the design works, how a payment flows, what we did to test it, and where it is still weak.""",
"""Small businesses pay contractors in stages: a deposit, a delivery, an acceptance test, a retention payment. Today that is trust, chat messages and bank transfers. Disputes are slow, and the buyer often pays before seeing the work.

AI agents could do this paperwork for us. But an agent that can move money is dangerous. It can be prompt-injected by a document, it can hallucinate an amount, and its key can leak. If the safety rules live in the agent's prompt, one clever message breaks them.

So the question for this topic is: how do you let an agent act, while the rules are enforced by something the agent cannot talk its way past?""",
"""The user owns the wallet. The agent only holds a session key. The contract decides.

The session key is scoped: it expires, it has a per-transaction cap and a daily cap, and it can pay only addresses on an allowlist. The owner key never goes to the agent. Anything outside those limits needs the user's own click.

We call the actions tiers. Tier 0, read. Tier 1, the agent proposes and a human approves. Tier 2, the agent acts alone, but only inside the on-chain caps. Tier 3, emergency: the agent, an approver or the owner can pause the wallet, and only the owner can unpause.

The point: a manipulated agent can be wrong, but it cannot be wrong with more money than the contract allows.""",
"""Here is the system. The user talks to an AI chat in the web app or in our desktop app. The agent has tools, and every tool call passes Guard 1, an off-chain policy engine that routes by tier, scans for injection and checks the allowlist.

Then Guard 2: the AgentWallet smart contract, one per user, which checks the same rules again on chain. A bug or jailbreak can get past Guard 1, but the contract still refuses.

Money in the real world is a simulated bank ledger in Postgres. When a deal is approved, the server debits the bank, mints a bank-backed token for exactly that amount, and locks it in an escrow contract. When the escrow pays out, the seller's bank is credited. The chain is a private Hyperledger Besu network. Only our app can write, but anyone can read it through a public explorer.""",
"""A buyer tells the agent: hire this designer, three milestones. The agent drafts the deal. Nothing is signed, and no money moves, until the buyer presses Approve and Lock.

Then the FileEscrow contract runs a fair exchange between strangers. The seller delivers an encrypted file, and its hash is recorded on chain. The buyer sees a watermarked preview. The key is released only after the buyer approves, or after the review window ends. Then the escrow pays the seller, and the seller's bank account is credited.

People without an account can join through a private link, protected by a one-time code and a password. A seller's payout destination is set once and then locked, so a leaked link cannot redirect money.""",
"""Here is the strongest evidence. We took the agent's own key and called a transfer to a random address, skipping the AI and the policy engine entirely. The contract reverted it with RecipientNotAllowed. This is a real transaction on our chain, and you can check it yourself with one command against our read-only RPC.

We also fixed an early bug where the agent could approve a file twice and skip the challenge window. Now the contract lets the agent approve only a delivered milestone.""",
"""The product is more than a contract. There is a desktop app for Windows and macOS, version 0.6.3, with its own signed chain of deals per device, an AI chat with web search, and MCP tools. There are company accounts where a group of signers controls the money, group payouts by vote, national-ID style KYC, top-ups by PromptPay and card in test mode, split payment, public job links where anyone can apply, and a neutral AI arbiter built from several models from different providers.

Everything runs on test money. No real funds or bank accounts are involved.""",
"""One team member took the role of security reviewer and attacked the system, rather than only reading it. The method: write a test that reproduces the bug first, then fix it.

Results: 115 contract tests including property tests that try every move on every state of an escrow; 15 web security tests; and a 22-step end-to-end smoke test that runs the main money flow on a throwaway stack. All of it runs in continuous integration on GitHub.

What it found: a High bug where a refund paid twice after a split ruling; a High bug where a company member who was not a signer could spend the company's money; and a silent arbiter who could lock money forever, for which we added a seven-day timeout. These fixes are in a pull request awaiting merge.

We also found that one desktop route could add a payee without a signer check. We found it in code review, not by exploiting it, and it is recorded as open.""",
"""We want to be honest about limits. In the demo, the server holds the owner keys. Production needs keys on the user's device or a multisig. The chain has one validator on a home computer. Production needs several, run by different parties. Disputes are switched off in the current build, so a buyer's protection against bad work is not approving, and a timeout that refunds a silent seller is the proposed fix. Desktop builds are unsigned, so updates need real code signing. All of these are written down in our security document as risks R1 to R15.""",
"""To recap: the agent proposes, the user approves, and the contract enforces. The rules do not live in a prompt. They live where the agent cannot argue with them. Everything you saw is on a public chain you can check yourself. Thank you. The app, the explorer and the evidence are linked on this slide."""
]


def build():
    prs = Presentation(); prs.slide_width = W; prs.slide_height = H
    blank = prs.slide_layouts[6]
    slides = []

    # 1: title
    s = prs.slides.add_slide(blank); slides.append(s)
    text(s, "KRAVII", .72, .72, 4.2, .55, 18, TEAL, True, "Aptos Display")
    text(s, "An AI agent that runs a wallet,\nand cannot lose your money", .72, 1.53, 10.8, 1.75, 41, NAVY, True, "Aptos Display")
    text(s, "TEAM KRAVII  ·  FINTECHATHON 2026  ·  TOPIC F", .75, 3.63, 7.0, .3, 15, MUTED, True)
    text(s, "formerly Agentic Wallet", .75, 4.05, 4.2, .3, 16, MUTED)
    box(s, 9.8, 1.18, 2.52, 3.64, PALE, PALE, True)
    text(s, "USER", 10.18, 1.72, 1.75, .28, 16, NAVY, True, align=PP_ALIGN.CENTER)
    text(s, "AGENT", 10.18, 2.73, 1.75, .28, 16, NAVY, True, align=PP_ALIGN.CENTER)
    text(s, "CONTRACT", 10.02, 3.76, 2.05, .28, 16, TEAL, True, align=PP_ALIGN.CENTER)
    connector(s, 11.06, 2.08, 11.06, 2.55); connector(s, 11.06, 3.08, 11.06, 3.58)
    footer(s, 1)

    # 2: problem
    s = prs.slides.add_slide(blank); slides.append(s); title(s, "The problem", "Why prompts are not a safety boundary")
    cards = [("STAGED PAYMENTS", "rely on trust", "01"), ("AI AGENTS", "can be fooled", "02"), ("PROMPT RULES", "break", "03")]
    for i, (h, b, n) in enumerate(cards):
        x = .82 + i * 4.16; box(s, x, 2.28, 3.56, 2.67, WHITE, LINE, True)
        text(s, n, x+.28, 2.58, .55, .35, 24, TEAL, True, "Aptos Display")
        text(s, h, x+.28, 3.20, 2.85, .30, 18, NAVY, True, "Aptos Display")
        text(s, b, x+.28, 3.74, 2.85, .42, 24, INK, True)
    text(s, "How can an agent act when its rules must survive manipulation?", .83, 5.65, 11.5, .4, 24, NAVY, True)
    footer(s, 2)

    # 3: authority
    s = prs.slides.add_slide(blank); slides.append(s); title(s, "User owns · agent holds a key · contract decides", "Three layers of authority", 30)
    for i, (a, b) in enumerate([("USER", "owns"), ("AGENT", "session key"), ("CONTRACT", "decides")]):
        y = 2.1 + i * 1.06; box(s, .82, y, 4.25, .72, PALE if i == 2 else WHITE, TEAL if i == 2 else LINE, True)
        text(s, a, 1.06, y+.18, 1.28, .22, 14, TEAL if i == 2 else NAVY, True)
        text(s, b, 2.35, y+.15, 2.35, .3, 20, INK, True)
        if i < 2: connector(s, 2.95, y+.73, 2.95, y+1.0)
    box(s, 5.65, 2.08, 6.72, 3.9, WHITE, LINE, True)
    label(s, "Action tiers", 5.95, 2.35, 2.0, .2)
    for i, (tier, action) in enumerate([("0", "Read"), ("1", "Propose"), ("2", "Auto within caps"), ("3", "Emergency pause")]):
        y = 2.83 + i*.66
        text(s, tier, 6.0, y, .35, .27, 16, TEAL, True, align=PP_ALIGN.CENTER)
        text(s, action, 6.65, y, 4.8, .27, 20, INK, True)
        if i < 3: connector(s, 6.0, y+.37, 6.0, y+.57, LINE, False)
    footer(s, 3)

    # 4: architecture
    s = prs.slides.add_slide(blank); slides.append(s); title(s, "Architecture", "Two guards; one enforcement boundary")
    nodes = [("User", .78, 2.1, 1.55), ("AI agent", 2.75, 2.1, 1.65), ("Guard 1\npolicy", 4.8, 2.1, 1.58), ("Guard 2\nAgentWallet", 6.82, 2.1, 1.85), ("FileEscrow", 9.15, 2.1, 1.72)]
    for i, (name, x, y, w) in enumerate(nodes):
        box(s, x, y, w, .88, PALE if i in (2,3) else WHITE, TEAL if i == 3 else LINE, True)
        text(s, name, x+.1, y+.18, w-.2, .48, 18, NAVY, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if i: connector(s, x-.38, y+.44, x-.06, y+.44)
    text(s, "CONTRACT BOUNDARY", 6.9, 1.58, 1.8, .2, 9, TEAL, True, align=PP_ALIGN.CENTER)
    box(s, .78, 4.2, 3.6, 1.0, WHITE, LINE, True); text(s, "Mock Bank ledger", 1.0, 4.45, 3.16, .3, 20, NAVY, True, align=PP_ALIGN.CENTER)
    box(s, 6.0, 4.2, 4.1, 1.0, PALE, TEAL, True); text(s, "Private Besu chain", 6.2, 4.34, 3.7, .3, 20, NAVY, True, align=PP_ALIGN.CENTER)
    text(s, "public read-only explorer", 6.2, 4.78, 3.7, .2, 12, MUTED, align=PP_ALIGN.CENTER)
    connector(s, 4.4, 4.7, 5.98, 4.7); text(s, "debit > mint > lock", 4.42, 4.3, 1.5, .22, 10, TEAL, True, align=PP_ALIGN.CENTER)
    connector(s, 7.75, 3.0, 7.75, 4.18)
    text(s, "When the escrow pays out, the seller's bank is credited. Only the app can write to the chain; anyone can read it.", .78, 5.7, 11.5, .7, 18, MUTED)
    footer(s, 4)

    # 5: flow
    s = prs.slides.add_slide(blank); slides.append(s); title(s, "A deal, step by step", "A six-step fair exchange")
    steps = ["Draft", "Approve\nand Lock", "Deliver\n(encrypted)", "Preview", "Approve", "Pay"]
    for i, step in enumerate(steps):
        x = .67 + i*2.07
        box(s, x, 2.6, 1.8, 1.5, PALE if i in (1,5) else WHITE, TEAL if i in (1,5) else LINE, True)
        text(s, str(i+1).zfill(2), x+.15, 2.78, .3, .18, 11, TEAL, True)
        text(s, step, x+.1, 3.1, 1.6, .8, 18, NAVY, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if i < 5: connector(s, x+1.83, 3.35, x+2.04, 3.35)
    text(s, "No money moves before Approve and Lock.", .67, 4.75, 9.0, .4, 24, NAVY, True)
    text(s, "The delivered file is encrypted and its hash is recorded on chain; the key is released only after approval.", .67, 5.35, 11.5, .7, 18, MUTED)
    footer(s, 5)

    # 6: proof
    s = prs.slides.add_slide(blank); slides.append(s); title(s, "Proof: the guard refuses the agent's own key", "An on-chain refusal")
    box(s, .8, 2.2, 5.12, 2.75, NAVY, NAVY, True)
    label(s, "Terminal", 1.12, 2.55, 1.4, .2, 11, rgb("6FE5DF"))
    text(s, "cast receipt 0x7f47a0c1...011f3191", 1.12, 3.1, 4.48, .42, 18, WHITE, True, "Courier New")
    text(s, "agentTransfer to a stranger reverted:", 1.12, 3.85, 4.25, .25, 16, rgb("C9D6DF"))
    text(s, "RecipientNotAllowed", 1.12, 4.22, 4.25, .28, 19, rgb("6FE5DF"), True, "Courier New")
    box(s, 6.45, 2.2, 5.98, 2.75, GREY, LINE, False)
    text(s, "[SCREENSHOT: explorer tx\n0x7f47a0c1...011f3191]", 6.82, 3.2, 5.25, .65, 17, MUTED, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    footer(s, 6)

    # 7: components
    s = prs.slides.add_slide(blank); slides.append(s); title(s, "Built around the core", "The product beyond the contract")
    labels = ["Desktop app\nWindows, macOS, 0.6.3", "Company accounts\nand group votes", "KYC\nPromptPay / card top-ups (test)", "Split payment", "Public job links", "Multi-model\nneutral arbiter"]
    for i, item in enumerate(labels):
        col, row = i % 3, i // 3; x, y = .85 + col*4.12, 2.15 + row*1.82
        box(s, x, y, 3.52, 1.32, WHITE, LINE, True)
        text(s, f"0{i+1}", x+.22, y+.22, .38, .2, 11, TEAL, True)
        text(s, item, x+.72, y+.30, 2.48, .72, 20, NAVY, True, valign=MSO_ANCHOR.MIDDLE)
    footer(s, 7)

    # 8: tests
    s = prs.slides.add_slide(blank); slides.append(s); title(s, "How we tested it", "Attack the system; reproduce the bug; fix it")
    numbers = [("115", "contract tests"), ("15", "web security tests"), ("22", "step smoke test")]
    for i, (num, label_text) in enumerate(numbers):
        y = 2.08 + i*1.1
        text(s, num, .9, y, 1.38, .54, 32, TEAL, True, "Aptos Display", align=PP_ALIGN.RIGHT)
        text(s, label_text, 2.45, y+.12, 2.25, .28, 20, NAVY, True)
    box(s, 5.45, 2.0, 6.78, 4.4, WHITE, LINE, True)
    label(s, "Found", 5.76, 2.32, 1.1, .2)
    findings = ["High — refund paid twice", "High — non-signer could spend company money", "Silent arbiter locks funds → 7-day timeout"]
    for i, finding in enumerate(findings):
        text(s, "•", 5.78, 2.85+i*.9, .2, .2, 18, TEAL, True)
        text(s, finding, 6.1, 2.84+i*.9, 5.6, .42, 20, INK, True)
    text(s, "Fixes in review (pull request awaiting merge)", 5.78, 5.8, 6.0, .3, 16, MUTED, True)
    footer(s, 8)

    # 9: risks
    s = prs.slides.add_slide(blank); slides.append(s); title(s, "Known risks, next steps", "What the demo does not claim")
    label(s, "Risk", .94, 2.08, 4.8, .22); label(s, "Next", 7.0, 2.08, 4.8, .22)
    risks = ["Keys held by server (demo)", "One validator", "Disputes off", "Unsigned desktop builds"]
    nxt = ["Device keys", "Multi-validator", "Code signing"]
    for i, value in enumerate(risks):
        y=2.65+i*.67; box(s, .82, y, 5.2, .48, PALE, PALE, True); text(s, value, 1.08, y+.09, 4.65, .25, 18, INK, True)
    for i, value in enumerate(nxt):
        y=2.65+i*.82; box(s, 6.82, y, 5.35, .58, WHITE, TEAL, True); text(s, value, 7.1, y+.13, 4.7, .28, 20, NAVY, True)
    footer(s, 9)

    # 10: summary
    s = prs.slides.add_slide(blank); slides.append(s)
    label(s, "Kravii", .72, .63, 2.0, .25)
    text(s, "The agent proposes,\nthe user approves,\nthe contract enforces.", .72, 1.28, 9.8, 1.72, 39, NAVY, True, "Aptos Display")
    line = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(.72), Inches(3.55), Inches(1.25), Inches(.06)); line.fill.solid(); line.fill.fore_color.rgb=TEAL; line.line.fill.background()
    links = [("APP", "https://agentic-wallet-inky.vercel.app"), ("EXPLORER", "https://agentic-wallet-inky.vercel.app/explorer"), ("EVIDENCE", "EVIDENCE.md")]
    for i, (kind, value) in enumerate(links):
        x=.75+i*4.14; label(s, kind, x, 4.35, 1.5, .18); text(s, value, x, 4.7, 3.85, .8, 16, NAVY, True)
    text(s, "Thank you.", .75, 5.85, 3.0, .4, 23, TEAL, True, "Aptos Display")
    footer(s, 10)

    for slide, speaker_note in zip(slides, NOTES):
        note(slide, speaker_note)
    prs.save(OUT)


if __name__ == "__main__":
    build()
